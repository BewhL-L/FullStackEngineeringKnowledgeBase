"""文档解析：支持 PDF / Word / TXT / Markdown / 图片（VLM 描述）"""
import os
import base64
from typing import Optional


def parse_image(file_path: str) -> str:
    """
    解析图片：用多模态大模型（VLM）生成图片内容描述，使其可被检索。
    优先用 VLM；失败时尝试本地 OCR。
    """
    from ..config import settings

    ext = os.path.splitext(file_path)[1].lower().lstrip(".")
    mime = {"jpg": "jpeg", "jpeg": "jpeg", "png": "png", "webp": "webp"}.get(
        ext, "png"
    )
    with open(file_path, "rb") as f:
        b64 = base64.b64encode(f.read()).decode()

    if settings.MULTIMODAL_ENABLED:
        try:
            from openai import OpenAI

            client = OpenAI(
                api_key=settings.LLM_API_KEY,
                base_url=settings.LLM_BASE_URL,
                timeout=settings.LLM_TIMEOUT,
            )
            resp = client.chat.completions.create(
                model=settings.LLM_MODEL,
                messages=[
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "text",
                                "text": "请描述这张图片的内容，如有文字请识别，如有表格请结构化输出。",
                            },
                            {
                                "type": "image_url",
                                "image_url": {
                                    "url": f"data:image/{mime};base64,{b64}"
                                },
                            },
                        ],
                    }
                ],
                max_tokens=600,
            )
            description = resp.choices[0].message.content
            if description:
                return f"[图片描述] {description}"
        except Exception:
            pass

    # VLM 失败，尝试本地 OCR
    try:
        from ..multimodal.parser import ocr_image
        text = ocr_image(file_path)
        if text.strip():
            return f"[图片OCR] {text}"
    except Exception:
        pass
    return "[图片] 无法识别图片内容"


def parse_pdf(file_path: str) -> str:
    """解析 PDF 文件，返回纯文本"""
    from PyPDF2 import PdfReader

    reader = PdfReader(file_path)
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        text_parts.append(page_text)
    return "\n".join(text_parts)


def parse_docx(file_path: str) -> str:
    """解析 Word（.docx）文件"""
    from docx import Document

    doc = Document(file_path)
    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
    # 表格内容也提取
    for table in doc.tables:
        for row in table.rows:
            row_text = " ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
    return "\n".join(paragraphs)


def parse_txt(file_path: str) -> str:
    """解析 TXT / Markdown 纯文本文件"""
    # 尝试 utf-8，失败则用 gbk
    for encoding in ["utf-8", "gbk", "latin-1"]:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    return ""


def parse_document(file_path: str) -> str:
    """
    统一文档解析入口
    根据文件扩展名自动选择解析器
    """
    ext = os.path.splitext(file_path)[1].lower()

    parsers = {
        ".pdf": parse_pdf,
        ".docx": parse_docx,
        ".doc": parse_docx,
        ".txt": parse_txt,
        ".md": parse_txt,
        ".markdown": parse_txt,
        ".jpg": parse_image,
        ".jpeg": parse_image,
        ".png": parse_image,
        ".webp": parse_image,
    }

    parser = parsers.get(ext)
    if parser is None:
        raise ValueError(f"不支持的文件类型：{ext}")

    text = parser(file_path)
    # 基础清洗：合并多余空白
    text = clean_text(text)
    return text


def clean_text(text: str) -> str:
    """文本清洗：去除多余空白但保留段落结构"""
    if not text:
        return ""
    # 统一换行符
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # 去除每行首尾空白
    lines = [line.strip() for line in text.split("\n")]
    # 合并 3 个以上连续空行为 2 个
    result = []
    empty_count = 0
    for line in lines:
        if not line:
            empty_count += 1
            if empty_count <= 2:
                result.append("")
        else:
            empty_count = 0
            result.append(line)
    return "\n".join(result).strip()
